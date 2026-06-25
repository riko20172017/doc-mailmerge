from openpyxl import load_workbook
from mailmerge import MailMerge, MailMergeOptions, OptionAutoUpdateFields, OptionKeepFields
import qrcode
from io import BytesIO
import tempfile
import os
from PIL import Image
from pathlib import Path
import warnings
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
import copy

# Вместо merge_templates попробуйте просто merge
with MailMerge('cover.docx') as doc:
    # Один диплом
    doc.merge(
        Фамилия="Иванов",
        Имя="Иван", 
        Отчество="Иванович",
        Оценки="5###4###3"
    )
    doc.write('test.docx')

# Проверяем
from docx import Document
doc = Document('test.docx')
for p in doc.paragraphs:
    if p.text.strip():
        print(p.text)